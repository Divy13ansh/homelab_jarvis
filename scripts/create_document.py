#!/usr/bin/env python3
import argparse
import pathlib
import sys


def convert(md_path: pathlib.Path, fmt: str, out: pathlib.Path | None) -> pathlib.Path:
    text = md_path.read_text(encoding="utf-8")
    if out is None:
        out = md_path.with_suffix(f".{fmt}")
    if fmt == "md":
        out.write_text(text, encoding="utf-8")
        return out
    if fmt == "pdf":
        try:
            import markdown
            from weasyprint import HTML
            html = markdown.markdown(text, extensions=["tables", "fenced_code"])
            styled = f"<html><head><meta charset='utf-8'><style>body{{font-family:sans-serif;max-width:800px;margin:40px auto;line-height:1.6}}pre{{background:#f4f4f4;padding:12px;overflow:auto}}table{{border-collapse:collapse;width:100%}}th,td{{border:1px solid #ddd;padding:8px}}a{{color:#0366d6}}</style></head><body>{html}</body></html>"
            HTML(string=styled).write_pdf(str(out))
            return out
        except ImportError:
            try:
                from reportlab.lib.pagesizes import A4
                from reportlab.lib.styles import getSampleStyleSheet
                from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
                styles = getSampleStyleSheet()
                doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
                story = []
                for line in text.splitlines():
                    if line.startswith("# "):
                        story.append(Paragraph(line[2:], styles["Heading1"]))
                    elif line.startswith("## "):
                        story.append(Paragraph(line[3:], styles["Heading2"]))
                    elif line.strip() == "":
                        story.append(Spacer(1, 8))
                    else:
                        story.append(Paragraph(line.replace("<", "&lt;").replace(">", "&gt;"), styles["BodyText"]))
                doc.build(story)
                return out
            except ImportError as e:
                print(f"pdf deps missing: {e}. Install: pip install weasyprint markdown reportlab", file=sys.stderr)
                sys.exit(2)
    if fmt == "docx":
        try:
            import markdown
            from docx import Document
            from docx.shared import Pt
            doc = Document()
            style = doc.styles["Normal"]
            style.font.size = Pt(10)
            html = markdown.markdown(text, extensions=["tables", "fenced_code"])
            for line in text.splitlines():
                if line.startswith("# "):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith("## "):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith("### "):
                    doc.add_heading(line[4:], level=3)
                elif line.strip() == "":
                    continue
                else:
                    doc.add_paragraph(line)
            doc.save(str(out))
            return out
        except ImportError as e:
            print(f"docx deps missing: {e}. Install: pip install python-docx markdown", file=sys.stderr)
            sys.exit(2)
    print(f"unknown format: {fmt}", file=sys.stderr)
    sys.exit(2)

def main():
    p = argparse.ArgumentParser(description="create_document helper — md → pdf/docx")
    p.add_argument("--input", required=True, help="input markdown path")
    p.add_argument("--format", required=True, choices=["md", "pdf", "docx"], help="output format")
    p.add_argument("--output", help="output path (default: input with new suffix)")
    p.add_argument("--title", help="optional title (unused, kept for interface compat)")
    args = p.parse_args()
    inp = pathlib.Path(args.input)
    if not inp.exists():
        print(f"input not found: {inp}", file=sys.stderr)
        sys.exit(1)
    out = pathlib.Path(args.output) if args.output else None
    result = convert(inp, args.format, out)
    print(str(result))

if __name__ == "__main__":
    main()
